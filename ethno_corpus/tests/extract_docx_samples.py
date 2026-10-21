#!/usr/bin/env python3
"""Извлекает текстовые фрагменты из всех .docx файлов для создания тестов"""

from pathlib import Path
from docx import Document as DocxDocument
import json

input_dir = Path("input")
output_file = Path("ethno_corpus/tests/docx_samples.json")

print(f"Scanning {input_dir}...")
print(f"Directories: {list(input_dir.iterdir())}\n")

samples = {}

# Ищем ВСЕ .docx файлы рекурсивно
docx_files = list(input_dir.rglob("*.docx"))
print(f"Found {len(docx_files)} .docx files\n")

for docx_file in sorted(docx_files):
    try:
        doc = DocxDocument(docx_file)

        # Извлекаем весь текст
        all_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                all_text.append(para.text.strip())

        text = "\n".join(all_text)

        if not text:
            print(f"[SKIP] Empty: {docx_file.name}")
            continue

        # Извлекаем до 10 непустых абзацев как фрагменты
        fragments = []
        for para in doc.paragraphs:
            p_text = para.text.strip()
            if p_text and len(p_text) > 5:
                # Ограничиваем размер до 150 символов для читаемости
                fragments.append(p_text[:150])
                if len(fragments) >= 10:  # До 10 фрагментов
                    break

        if fragments:
            samples[docx_file.name] = {
                "path": str(docx_file),
                "rel_path": str(docx_file.relative_to(input_dir)),
                "text_length": len(text),
                "fragments": fragments
            }
            print(f"[OK] {docx_file.relative_to(input_dir)}: {len(text)} chars, {len(fragments)} fragments")

    except Exception as e:
        print(f"[ERROR] {docx_file.name}: {e}")

# Сохраняем в JSON
output_file.write_text(json.dumps(samples, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"\nSaved {len(samples)} files to {output_file}")
